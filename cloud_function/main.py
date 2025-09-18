import functions_framework
from google.cloud import storage, firestore
import tempfile, os, uuid, sys
import fitz, pdfplumber
from docx import Document
from docx.shared import Inches
import docx
import google.generativeai as genai
import json
# --- CONFIG ---
API_KEY = "API KEY"
MODEL_NAME = "gemini-2.5-pro"
BUCKET_NAME = "aianalystingestion"

# Ensure UTF-8 printing
sys.stdout.reconfigure(encoding='utf-8')

# Google clients
storage_client = storage.Client()
fs_client = firestore.Client()

# Gemini config
genai.configure(api_key=API_KEY)

# ---------- DOCX Generation ----------
def gemini_ai(sometext, imgpath=None):
    m = genai.GenerativeModel("gemini-2.0-flash")
    try:
        if imgpath:
            res = m.generate_content([sometext, genai.upload_file(imgpath)])
        else:
            res = m.generate_content(sometext)
        if res and getattr(res, "text", None):
            return res.text
        else:
            return "no ai answer"
    except Exception as er:
        return f"error {er}"

def cleanline(txt):
    return txt.replace("*", "").replace("#", "").replace("-", "").replace("•", "").strip()

def put_ai_text(docfile, txt):
    lines = txt.split("\n")
    now = None
    for l in lines:
        l = cleanline(l)
        if not l:
            continue
        low = l.lower()
        if low.startswith("context"):
            docfile.add_heading("Context", level=3)
            now = "context"
        elif low.startswith("key points"):
            docfile.add_heading("Key Points", level=3)
            now = "points"
        elif low.startswith("implications"):
            docfile.add_heading("Implications", level=3)
            now = "imp"
        else:
            if now == "points":
                docfile.add_paragraph(l, style="List Bullet")
            else:
                docfile.add_paragraph(l)

def makeReport(pdfname, outname="ai_output.docx"):
    d = Document()
    p = fitz.open(pdfname)
    made_imgs = []
    with pdfplumber.open(pdfname) as pp:
        for pageNo in range(len(p)):
            pg = p[pageNo]
            plpg = pp.pages[pageNo]
            theText = pg.get_text("text")

            d.add_heading("Page " + str(pageNo+1), level=1)

            if theText.strip():
                d.add_heading("Original Text:", level=2)
                d.add_paragraph(theText)
                ans = gemini_ai(
                    f"""Analyze this text from page {pageNo+1}.

                    Context: say what it is about.
                    Key Points: main details and insights.
                    Implications: meaning or importance.

                    Text:
                    {theText}"""
                )
                d.add_heading("Detailed Explanation (AI):", level=2)
                put_ai_text(d, ans)

            imgs = pg.get_images(full=True)
            for i, im in enumerate(imgs):
                xref = im[0]
                base = p.extract_image(xref)
                data = base["image"]
                ext = base["ext"]
                fname = f"/tmp/pg{pageNo+1}_img{i}.{ext}"
                with open(fname, "wb") as f:
                    f.write(data)
                made_imgs.append(fname)
                d.add_heading("Image " + str(i+1) + " from Page " + str(pageNo+1), level=2)
                try:
                    d.add_picture(fname, width=Inches(5))
                except:
                    pass
                imans = gemini_ai(
                    f"""Analyze this image/graph from page {pageNo+1}.

                    Context: what is shown.
                    Key Points: important things or patterns.
                    Implications: why it matters.
                    """,
                    imgpath=fname
                )
                d.add_heading("Image Analysis (AI):", level=2)
                put_ai_text(d, imans)

            tbls = plpg.extract_tables()
            for tnum, t in enumerate(tbls):
                d.add_heading("Table " + str(tnum+1) + " from Page " + str(pageNo+1), level=2)
                tb = d.add_table(rows=len(t), cols=len(t[0]))
                for r, row in enumerate(t):
                    for c, cell in enumerate(row):
                        tb.cell(r, c).text = str(cell) if cell else ""
                tabtxt = "\n".join([", ".join(r) for r in t])
                tabsum = gemini_ai(
                    f"""Analyze this table from page {pageNo+1}.

                    Context: what the table is about.
                    Key Points: main values and findings.
                    Implications: meaning or impact.

                    Table:
                    {tabtxt}"""
                )
                d.add_heading("Table Analysis (AI):", level=2)
                put_ai_text(d, tabsum)
    d.save(outname)
    for f in made_imgs:
        try:
            os.remove(f)
        except:
            pass
    print("AI report saved to", outname)


# ---------- DOCX → TXT Structured Output ----------
def extract_text_from_docx(path):
    doc = docx.Document(path)
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

def process_with_gemini_txt(text):
    prompt = f"""
    Extract structured company details from the docx in this format(all info to be covered):

    company_name -> <company name>
    description -> <short description of the company>
    website -> <url>
    founded_year -> <year if available, else None>
    head_quarters -> <location>
    employee_count -> <integer if available, else None>
    industry -> <industry>
    business_model -> <business model>

    Text to analyze:
    {text}
    """
    model = genai.GenerativeModel(MODEL_NAME)
    response = model.generate_content(prompt)
    return response.text

def save_to_txt(content, filename):
    with open(filename, "w", encoding="utf-8") as f:
        f.write(content)


# ---------- MAIN CLOUD FUNCTION ----------
@functions_framework.cloud_event
def process_pdf(cloud_event):
    data = cloud_event.data
    bucket_name = data.get("bucket", BUCKET_NAME)
    file_name = data.get("name")

    if not file_name or not file_name.endswith(".pdf"):
        print("Skipping non-pdf file:", file_name)
        return

    # Use bucket folder name as company_id to keep consistent
    company_id = file_name.split("/")[0]
    founder_id = company_id  # same as company_id

    bucket = storage_client.bucket(bucket_name)
    blob = bucket.blob(file_name)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        blob.download_to_filename(tmp.name)
        pdf_path = tmp.name

    # 1️⃣ Generate DOCX
    outdoc = f"/tmp/{company_id}_ai_output.docx"
    makeReport(pdf_path, outdoc)

    # Upload DOCX to GCS
    report_blob = bucket.blob(f"{company_id}/ai_output.docx")
    report_blob.upload_from_filename(outdoc)

    # Save DOCX reference in Firestore
    fs_client.collection("reports").document(company_id).set({
        "docx_path": f"gs://{bucket_name}/{company_id}/ai_output.docx",
        "company_id": company_id
    })

    # 2️⃣ Extract text from DOCX
    extracted_text = extract_text_from_docx(outdoc)
    print("🔹 Extracted Text Preview:\n", extracted_text[:500])

    # # 3️⃣ Generate structured text file
    # result = process_with_gemini_txt(extracted_text, company_id, founder_id)
    # print("🔹 Gemini Raw Output:\n", result)

    # outtxt = f"/tmp/{company_id}_company_founder_data.txt"
    # save_to_txt(result, outtxt)

    # # Upload TXT to GCS
    # txt_blob = bucket.blob(f"{company_id}/company_founder_data.txt")
    # txt_blob.upload_from_filename(outtxt)

    # print(f"[✅] Saved text to gs://{bucket_name}/{company_id}/company_founder_data.txt")

    # # Cleanup
    # try:
    #     os.remove(pdf_path)
    #     os.remove(outdoc)
    #     os.remove(outtxt)
    # except:
    #     pass
    # 3️⃣ Generate structured text file
    result = process_with_gemini_txt(extracted_text)
    print("🔹 Gemini Raw Output:\n", result)

    outtxt = f"/tmp/{company_id}_company_founder_data.txt"
    save_to_txt(result, outtxt)

    # Upload TXT to GCS
    txt_blob = bucket.blob(f"{company_id}/company_founder_data.txt")
    txt_blob.upload_from_filename(outtxt)

    # 4️⃣ Save the same content as company.json
    outjson = f"/tmp/{company_id}_company.json"

    # clean result (remove possible ```json or ``` markers)
    clean_result = result.strip()
    if clean_result.startswith("```json"):
        clean_result = clean_result[len("```json"):].strip()
    if clean_result.startswith("```"):
        clean_result = clean_result[len("```"):].strip()
    if clean_result.endswith("```"):
        clean_result = clean_result[:-3].strip()

    # ensure it's valid JSON
    data = json.loads(clean_result)

    # write clean JSON
    with open(outjson, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4)

    # upload
    json_blob = bucket.blob(f"{company_id}/company.json")
    json_blob.upload_from_filename(outjson)

    print(f"[✅] Saved text to gs://{bucket_name}/{company_id}/company_founder_data.txt")
    print(f"[✅] Saved JSON to gs://{bucket_name}/{company_id}/company.json")

    # Cleanup
    try:
        os.remove(pdf_path)
        os.remove(outdoc)
        os.remove(outtxt)
        os.remove(outjson)
    except:
        pass
